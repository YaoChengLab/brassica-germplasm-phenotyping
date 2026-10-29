"""Compare regenerated outputs with archived, rounded model results."""
from pathlib import Path
import json,sys
import numpy as np
import pandas as pd
from openpyxl import load_workbook
root=Path(__file__).resolve().parents[1];out=Path(sys.argv[1]);checks=[];precision_differences=[]
def record(name,ok,detail=''):
    checks.append({'check':name,'passed':bool(ok),'detail':detail})
for filename in ['Table_model_comparison.csv','Table_association.csv','Table_coefficients.csv','Table_batch_diagnostics.csv']:
    old=pd.read_csv(root/'reference'/filename);new=pd.read_csv(out/'all_batches'/filename)
    ok=old.shape==new.shape and list(old.columns)==list(new.columns)
    for col in old.columns:
        if filename == 'Table_association.csv' and col in ['within_batch_spearman_rho','p_spearman']:
            # Rounding creates/removes ties. Record every change separately;
            # these are not claims of exact reproduction of rank statistics.
            limit = 0.001 if col == 'within_batch_spearman_rho' else 0.01
            ok &= np.allclose(old[col],new[col],rtol=0,atol=limit)
            for i in range(len(old)):
                if old.loc[i,col] != new.loc[i,col]:
                    precision_differences.append({'trait':old.loc[i,'trait'],'statistic':col,'archived':float(old.loc[i,col]),'rerun':float(new.loc[i,col]),'reason':'Rank ties can change when archived means are rounded.'})
        elif pd.api.types.is_numeric_dtype(old[col]):
            ok &= np.allclose(old[col],new[col],rtol=6e-4,atol=1e-7,equal_nan=True)
        else:ok &= old[col].fillna('').equals(new[col].fillna(''))
    record('Archived '+filename,ok,'Relative tolerance 0.0006 for four-significant-digit outputs; absolute tolerance 1e-7. Rank diagnostics separately allow absolute changes of 0.001 in rho and 0.01 in p; all changes are listed.')
old=pd.read_csv(root/'reference/Table_S7_original.csv').set_index('Statistic')
new=pd.read_csv(out/'Table_S7.csv',dtype=str)
keys=['Analysis units (accession × batch)','Accessions / batches','Partial r, ln PLA vs ln FW (95% CI)',
'M1 cross-validated R² (mean ± SD, 50 repeats)','M2 cross-validated R² (mean ± SD)',
'M2 vs M1, nested F test','M2 vs M1, cluster-robust Wald test, P']
for i,key in enumerate(keys):
    for j,col in enumerate(old.columns,1):
        expected=str(old.loc[key,col])
        if i==1:expected=expected.split(' / ')[0]
        if i==5:expected=expected.split('P = ')[-1]
        record(f'S7 row {i+1}, analysis {j}',new.iloc[i,j]==expected,f'Expected {expected}; obtained {new.iloc[i,j]}')
for case,n,na,nb in [('all_batches',269,257,12),('without_B12_B13',221,211,10)]:
    summary=json.loads((out/case/'summary.json').read_text(encoding='utf-8'))
    record(case+' sample counts',(summary['n_units'],summary['n_accessions'],summary['n_batches'])==(n,na,nb))
    folds=pd.read_csv(out/case/'CV_folds.csv')
    record(case+' accession-grouped CV',len(folds)==n*50 and folds.groupby(['repeat','accession']).fold.nunique().eq(1).all())
    record(case+' 50 repeats and 10 folds',folds.repeat.nunique()==50 and folds.groupby('repeat').fold.nunique().eq(10).all())
    coords=pd.read_csv(out/case/'Figure7_plot_data.csv')
    record(case+' plot data',len(coords)==n and np.isfinite(coords.select_dtypes(include='number')).all().all())
    metrics=pd.read_csv(out/case/'CV_repeat_metrics.csv')
    record(case+' four models x 50 repeats',len(metrics)==200 and metrics.groupby('model').size().eq(50).all())
wb=load_workbook(out/'Table_S7.xlsx',data_only=True)
record('Excel and CSV match',[[str(c.value) for c in row] for row in wb['S7'].iter_rows(min_row=5,max_row=11,max_col=3)]==new.values.tolist())
result={'passed':all(x['passed'] for x in checks),'checks':checks,'documented_precision_differences':precision_differences,
'precision_note':'Models rerun from archived means rounded to six significant digits; agreement is checked at reported precision, not byte identity with the original figure.',
'raw_weight_limit':'Raw plant-level fresh-weight records are unavailable; their exclusion and averaging steps are outside this package.'}
(out/'verification.json').write_text(json.dumps(result,indent=2)+'\n',encoding='utf-8')
if not result['passed']:
    for x in checks:
        if not x['passed']:print('FAILED:',x)
    raise SystemExit(1)
print(f'PASS: {len(checks)} checks; {len(precision_differences)} rank-statistic precision differences documented')
